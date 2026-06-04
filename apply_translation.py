#!/usr/bin/env python3
"""Apply a per-language translation map to slides.pptx -> slides_<lang>.pptx.

Replaces run text by EXACT match against the map keys, so colored markers,
code, product names, ports, URLs and citations (not in the map) stay verbatim.
"""
import json, sys, shutil
from pptx import Presentation

lang = sys.argv[1]                      # e.g. "es"
map_path = sys.argv[2]                  # e.g. /tmp/trans_es.json
src = "slides.pptx"
dst = f"slides_{lang}.pptx"

with open(map_path, encoding="utf-8") as f:
    tmap = json.load(f)

shutil.copyfile(src, dst)
prs = Presentation(dst)

replaced = 0
used_keys = set()
for slide in prs.slides:
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        for para in shp.text_frame.paragraphs:
            for r in para.runs:
                if r.text in tmap:
                    new = tmap[r.text]
                    used_keys.add(r.text)
                    if new != r.text:
                        r.text = new
                        replaced += 1
prs.save(dst)

# coverage vs the canonical to-translate list
with open("/tmp/strings_to_translate.json", encoding="utf-8") as f:
    canon = json.load(f)
missing = [s for s in canon if s not in tmap]

print(f"[{lang}] saved {dst} | runs changed={replaced} | map keys={len(tmap)} "
      f"| canon={len(canon)} | missing-from-map={len(missing)}")
if missing:
    print(f"[{lang}] FIRST MISSING (left English):")
    for s in missing[:15]:
        print("   -", repr(s))
