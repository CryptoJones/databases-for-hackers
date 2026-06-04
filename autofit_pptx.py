#!/usr/bin/env python3
"""Shrink-to-fit text in a .pptx so translated (longer) strings stop overflowing.

The decks lay text out with ABSOLUTE positions and FIXED font sizes tuned for
English. Longer translations (ES/FR/DE/RU) wrap to extra lines and collide with
the box below ("top and bottom" overlap). This pass measures each text box in
the font LibreOffice actually renders (Calibri -> Noto Sans substitute) at the
real export DPI and shrinks the font until the wrapped text fits the *gap to the
next box*, not just the box's own slack height. Adjacent full-width boxes of the
same size (a hero sentence split across boxes) are scaled by one uniform factor
so they stay consistent.

Usage:  autofit_pptx.py in.pptx out.pptx [--report]
"""
import sys, unicodedata
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from PIL import ImageFont

EMU = 914400
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
DPI = 1920 / SLIDE_W_IN          # == 144.0 ; matches pdftoppm -scale-to-x 1920
PT2PX = DPI / 72.0               # == 2.0
CAP = 0.85                       # visual height of ONE line / em (ink, not advance)
ADVANCE = 1.25                   # line-to-line baseline advance / em
LINE_FACTOR = ADVANCE            # (kept for cluster line-height estimate)
SAFETY = 0.95                    # leave headroom for metric mismatch
MIN_SCALE = 0.34                 # never shrink below this fraction of original
WIDE_IN = 9.0                    # boxes >= this wide can form hero clusters
M_SLIDE = 0.12                   # keep text this far inside the slide's left/right edges (in)

REG = "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"
BOLD = "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"
_cache = {}


def _font(px, bold):
    px = max(1, int(round(px)))
    key = (px, bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(BOLD if bold else REG, px)
    return _cache[key]


def _is_cjk(ch):
    o = ord(ch)
    return (0x3000 <= o <= 0x9FFF or 0xAC00 <= o <= 0xD7A3 or
            0xF900 <= o <= 0xFAFF or 0xFF00 <= o <= 0xFFEF)


def text_width(s, px, bold):
    """Pixel width of a single run of text; CJK glyphs (no Noto Sans coverage)
    approximated at ~1 em each since they render in a CJK fallback face."""
    f = _font(px, bold)
    w = 0.0
    run = []
    for ch in s:
        if _is_cjk(ch):
            if run:
                w += f.getlength("".join(run)); run = []
            w += px  # ~1 em
        else:
            run.append(ch)
    if run:
        w += f.getlength("".join(run))
    return w


def wrap_count(text, px, bold, max_w):
    """Number of wrapped lines for text at given size within max_w pixels.
    Breaks on spaces; over-long tokens (e.g. CJK) break per-character."""
    lines = 0
    for hard in text.split("\n"):
        tokens = hard.split(" ")
        cur = ""
        for tok in tokens:
            trial = tok if not cur else cur + " " + tok
            if text_width(trial, px, bold) <= max_w or not cur:
                cur = trial
            else:
                lines += 1
                cur = tok
            # break a single token that is itself too wide (CJK / long word)
            while text_width(cur, px, bold) > max_w and len(cur) > 1:
                lo, hi = 1, len(cur)
                while lo < hi:
                    mid = (lo + hi + 1) // 2
                    if text_width(cur[:mid], px, bold) <= max_w:
                        lo = mid
                    else:
                        hi = mid - 1
                lines += 1
                cur = cur[lo:]
        lines += 1
    return max(1, lines)


def para_info(p):
    """Return (text, max_pt, any_bold, line_spacing) for a paragraph, or None."""
    txt = "".join(r.text for r in p.runs)
    if not txt.strip():
        return None
    sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
    if not sizes:
        return None
    bold = any(bool(r.font.bold) for r in p.runs)
    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
    return txt, max(sizes), bold, ls


def frame_paras(shape):
    out = []
    for p in shape.text_frame.paragraphs:
        info = para_info(p)
        if info:
            out.append(info)
    return out


def block_height_px(paras, scale, usable_w_px):
    """Total rendered text height. The FIRST line of the frame occupies ~CAP em
    of ink; every additional line/paragraph adds ~ADVANCE em (x line_spacing).
    Modeling the first line as CAP (not ADVANCE) is what stops single tall lines
    from being treated as overflowing their gap to the next box."""
    total = 0.0
    first_pt = paras[0][1] * scale * PT2PX
    for txt, pt, bold, ls in paras:
        px = pt * scale * PT2PX
        n = wrap_count(txt, px, bold, usable_w_px)
        total += n * px * ADVANCE * ls
    total -= (ADVANCE - CAP) * first_pt   # relieve the first line
    return max(first_pt * CAP, total)


def fits(paras, scale, usable_w_px, avail_h_px):
    # width: every paragraph's longest unbroken need must wrap within usable_w
    for txt, pt, bold, ls in paras:
        px = pt * scale * PT2PX
        # a single token wider than usable (after char-break it's fine, handled in wrap)
        if wrap_count(txt, px, bold, usable_w_px) < 1:
            return False
    return block_height_px(paras, scale, usable_w_px) <= avail_h_px * SAFETY


def geom(sh):
    if None in (sh.left, sh.top, sh.width, sh.height):
        return None
    return (sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)


def available_height(box, boxes):
    l, t, w, h = box
    avail = min(h, SLIDE_H_IN - t)
    for (ol, ot, ow, oh) in boxes:
        if ot <= t + 0.02:
            continue
        if ol < l + w and ol + ow > l:          # horizontal overlap
            avail = min(avail, ot - t)
    return max(0.1, avail)


def fit_scale(paras, box, boxes):
    l, t, w, h = box
    ml = mr = 0.1
    usable_w_px = max(10.0, (w - ml - mr) * DPI)
    avail_h_px = available_height(box, boxes) * DPI
    if fits(paras, 1.0, usable_w_px, avail_h_px):
        return 1.0
    s = 1.0
    while s > MIN_SCALE:
        s -= 0.02
        if fits(paras, s, usable_w_px, avail_h_px):
            return round(s, 3)
    return MIN_SCALE


def hscale_cap(shape, geom):
    """Horizontal-fit cap for boxes that DON'T soft-wrap (auto_size=SHAPE_TO_FIT_TEXT
    or word_wrap=False). Such a shape grows to fit its text on one line per paragraph,
    so a translated (longer) line can run off the slide's left/right edge and clip.
    Return the largest scale (<=1.0) at which every hard line stays inside the slide
    given the box position and paragraph alignment; 1.0 when not applicable / already
    inside. The vertical-gap pass handles everything else."""
    tf = shape.text_frame
    try:
        autosz = tf.auto_size
    except Exception:
        autosz = None
    nowrap = (autosz == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT) or (tf.word_wrap is False)
    if not nowrap:
        return 1.0
    l, t, w, h = geom
    # Only boxes whose FRAME pokes past the visible slide edge can clip at that edge in a
    # way font-shrink fixes (an in-bounds box wraps its text inside the slide). This is the
    # slide-2 bio pattern: wide SHAPE_TO_FIT boxes placed with negative left.
    if l >= M_SLIDE and (l + w) <= SLIDE_W_IN - M_SLIDE:
        return 1.0
    cx = l + w / 2.0
    ml = (tf.margin_left / EMU) if tf.margin_left is not None else 0.1
    mr = (tf.margin_right / EMU) if tf.margin_right is not None else 0.1
    cap = 1.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if not sizes:
            continue
        pt = max(sizes)
        bold = any(bool(r.font.bold) for r in p.runs)
        nw_in = text_width(txt, pt * PT2PX, bold) / DPI       # natural one-line width (in)
        if nw_in <= 0:
            continue
        align = p.alignment
        if align == PP_ALIGN.CENTER:
            allowed = 2.0 * min(cx - M_SLIDE, SLIDE_W_IN - M_SLIDE - cx)
        elif align == PP_ALIGN.RIGHT:
            allowed = (l + w - mr) - M_SLIDE
        else:                                                 # left / inherited
            allowed = (SLIDE_W_IN - M_SLIDE) - (l + ml)
        allowed = max(0.3, allowed)
        if nw_in > allowed:
            cap = min(cap, allowed / nw_in)
    return round(max(cap, MIN_SCALE), 3)


UNDERLINE_H = 0.15      # a decorative underline is a thin (<= this, in) empty bar
UNDERLINE_NEAR = 0.30   # ...sitting within this far (in) of the title box bottom
UNDERLINE_SAFETY = 0.92 # cap to this fraction of usable width so the one-line title
                        # clears the box with slack (our text_width vs LibreOffice's
                        # metrics differ by a few %; a boundary-tight cap still wraps)


def underline_title_cap(shape, geom, shapes):
    """Stop a SHAPE_TO_FIT_TEXT title from wrapping onto its decorative underline.

    These titles grow DOWNWARD to fit wrapped text (auto_size = shape-to-fit). A
    thin empty bar sits flush at the title box's bottom as a design underline; the
    English title is one line ending just above it. A longer translation wraps to a
    2nd line that renders ON the bar (and under the warning icon). The vertical-gap
    pass can't see it: the bar is an empty, non-text shape, and the box's own height
    already 'contains' two lines. So: when such an underline sits at the box bottom,
    cap the font so every paragraph fits on ONE line within the box's own width."""
    tf = shape.text_frame
    try:
        autosz = tf.auto_size
    except Exception:
        return 1.0
    if autosz != MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT:
        return 1.0
    l, t, w, h = geom
    bottom = t + h
    has_underline = False
    for s in shapes:
        if s is shape:
            continue
        try:
            if None in (s.left, s.top, s.width, s.height):
                continue
            sl, st, sw, sh2 = (s.left / EMU, s.top / EMU,
                               s.width / EMU, s.height / EMU)
        except Exception:
            continue
        if sh2 > UNDERLINE_H:                       # not a thin bar
            continue
        if s.has_text_frame and "".join(
                r.text for p in s.text_frame.paragraphs for r in p.runs).strip():
            continue                                # has text -> not a bare underline
        if not (bottom - UNDERLINE_NEAR <= st <= bottom + UNDERLINE_NEAR):
            continue                                # not flush with title bottom
        if sl < l + w and sl + sw > l:              # horizontal overlap
            has_underline = True
            break
    if not has_underline:
        return 1.0
    ml = (tf.margin_left / EMU) if tf.margin_left is not None else 0.1
    mr = (tf.margin_right / EMU) if tf.margin_right is not None else 0.1
    usable = max(0.5, w - ml - mr)
    cap = 1.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if not sizes:
            continue
        pt = max(sizes)
        bold = any(bool(r.font.bold) for r in p.runs)
        nw_in = text_width(txt, pt * PT2PX, bold) / DPI   # natural one-line width (in)
        target = usable * UNDERLINE_SAFETY
        if nw_in > target:
            cap = min(cap, target / nw_in)
    return round(max(cap, MIN_SCALE), 3)


SUBTITLE_HERO_PT = 30.0 # only a large 'hero' title (>= this pt) is targeted; body /
                        # bullet boxes (<= 24pt here) are meant to wrap and are left alone
SUBTITLE_SAFETY = 0.92  # cap a one-line title to this fraction of usable width (same
                        # metric-mismatch headroom as UNDERLINE_SAFETY)


def subtitle_title_cap(shape, geom, shapes):
    """Stop a SHAPE_TO_FIT_TEXT 'hero' title from wrapping onto a subtitle below it.

    Sibling of underline_title_cap, for a different obstacle. A large centered hero
    title grows DOWNWARD (auto_size = shape-to-fit). The English title is one line; a
    longer translation whose natural one-line width is ~99% of the box width wraps its
    last word/glyphs onto a 2nd line that collides with a NON-EMPTY subtitle box placed
    just below. The vertical-gap pass misses it: our width metric under-counts CJK by a
    few % and predicts one line, while LibreOffice wraps.

    To stay narrow (a 'box below' alone matches almost every stacked layout) we fire
    only on the structural anomaly that defines this pattern: the title is a LARGE font
    (>= SUBTITLE_HERO_PT) AND a non-empty text box's top sits ABOVE the title box's
    authored bottom -- i.e. it INTRUDES into the title's own box, which the designer only
    did because the title was assumed to be one line. Normal stacked content sits below
    the title bottom and is ignored. Then cap so every paragraph fits one line in width."""
    tf = shape.text_frame
    try:
        autosz = tf.auto_size
    except Exception:
        return 1.0
    if autosz != MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT:
        return 1.0
    title_sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs
                   if r.font.size is not None and "".join(rr.text for rr in p.runs).strip()]
    if not title_sizes or max(title_sizes) < SUBTITLE_HERO_PT:
        return 1.0                                   # not a hero-size title
    l, t, w, h = geom
    bottom = t + h
    has_below = False
    for s in shapes:
        if s is shape:
            continue
        try:
            if None in (s.left, s.top, s.width, s.height):
                continue
            sl, st, sw, sh2 = (s.left / EMU, s.top / EMU,
                               s.width / EMU, s.height / EMU)
        except Exception:
            continue
        if not s.has_text_frame:
            continue
        if not "".join(r.text for p in s.text_frame.paragraphs
                       for r in p.runs).strip():
            continue                                 # empty (e.g. underline bar) -> ignore
        if not (t + 0.10 < st < bottom):
            continue                                 # must INTRUDE into the title's box
        if sl < l + w and sl + sw > l:               # horizontal overlap
            has_below = True
            break
    if not has_below:
        return 1.0
    ml = (tf.margin_left / EMU) if tf.margin_left is not None else 0.1
    mr = (tf.margin_right / EMU) if tf.margin_right is not None else 0.1
    usable = max(0.5, w - ml - mr)
    usable_px = usable * DPI
    cap = 1.0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not txt.strip():
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if not sizes:
            continue
        pt = max(sizes)
        bold = any(bool(r.font.bold) for r in p.runs)
        # Only the MARGINAL wrap is ours: the model still thinks the line fits (1 line)
        # but its natural width is so close to the box that LibreOffice's slightly wider
        # metrics wrap the tail onto the intruding subtitle. When the model already wraps
        # (>=2 lines), fit_scale shrinks it to a clean multi-line title in the gap -> leave it.
        if wrap_count(txt, pt * PT2PX, bold, usable_px) != 1:
            continue
        nw_in = text_width(txt, pt * PT2PX, bold) / DPI   # natural one-line width (in)
        target = usable * SUBTITLE_SAFETY
        if nw_in > target:
            cap = min(cap, target / nw_in)
    return round(max(cap, MIN_SCALE), 3)


def process_slide(slide, report):
    text_shapes = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        g = geom(sh)
        if g is None:
            continue
        paras = frame_paras(sh)
        if not paras:
            continue
        text_shapes.append((sh, g, paras))
    boxes = [g for (_, g, _) in text_shapes]

    # per-box scale: vertical gap fit, horizontal slide-edge cap for off-slide
    # non-wrapping boxes, and single-line cap for titles sitting on a decorative underline
    all_shapes = list(slide.shapes)
    scales = {}
    for i, (sh, g, paras) in enumerate(text_shapes):
        scales[i] = min(fit_scale(paras, g, boxes), hscale_cap(sh, g),
                        underline_title_cap(sh, g, all_shapes),
                        subtitle_title_cap(sh, g, all_shapes))

    # uniform scaling for adjacent, same-size, full-width hero clusters
    wide = [i for i, (sh, g, paras) in enumerate(text_shapes)
            if g[2] >= WIDE_IN]
    wide.sort(key=lambda i: text_shapes[i][1][1])  # by top
    used = set()
    for a in wide:
        if a in used:
            continue
        _, (la, ta, wa, ha), pa = text_shapes[a]
        size_a = max(pt for _, pt, _, _ in pa)
        cluster = [a]
        cur_top = ta
        cur_size = size_a
        for b in wide:
            if b <= a or b in used:
                continue
            _, (lb, tb, wb, hb), pb = text_shapes[b]
            size_b = max(pt for _, pt, _, _ in pb)
            line_h = cur_size * LINE_FACTOR / 72.0
            if abs(size_b - cur_size) < 0.6 and (tb - cur_top) <= line_h * 1.6:
                cluster.append(b)
                cur_top = tb
        if len(cluster) > 1:
            smin = min(scales[i] for i in cluster)
            for i in cluster:
                scales[i] = smin
                used.add(i)

    changed = []
    for i, (sh, g, paras) in enumerate(text_shapes):
        s = scales[i]
        if s >= 0.999:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    newpt = round(r.font.size.pt * s * 2) / 2.0
                    r.font.size = Pt(max(6.0, newpt))
        changed.append((g, s, paras[0][0][:40]))
    return changed


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src, dst = sys.argv[1], sys.argv[2]
    report = "--report" in sys.argv
    prs = Presentation(src)
    total = 0
    for n, slide in enumerate(prs.slides, 1):
        ch = process_slide(slide, report)
        total += len(ch)
        if report and ch:
            print(f"slide {n}: {len(ch)} box(es) shrunk")
            for g, s, t in ch:
                print(f"    x{s:.2f}  {t!r}")
    prs.save(dst)
    print(f"[autofit] {src} -> {dst}: {total} boxes shrunk")


if __name__ == "__main__":
    main()
